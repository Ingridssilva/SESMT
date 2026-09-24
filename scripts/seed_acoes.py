"""
seed_acoes.py — Importa dados da planilha Acompanhamento_ações_SESMT_2026.xlsx
para a tabela acoes_sesmt no banco de dados.

Uso:
    DATABASE_URL=postgresql://... python scripts/seed_acoes.py
"""

import os
import sys
from datetime import date

import pg8000.native as pg
from openpyxl import load_workbook

XLSX_PATH = os.path.join(os.path.dirname(__file__), '..', 'data', 'Acompanhamento_ações_SESMT_2026.xlsx')

def get_conn():
    url = os.environ.get('DATABASE_URL', 'postgresql://postgres:postgres@localhost:5432/sesmt')
    url = url.replace('postgresql://', '').replace('postgres://', '')
    if '?' in url:
        url, params = url.split('?', 1)
    else:
        params = ''
    auth, rest = url.split('@', 1)
    user, password = auth.split(':', 1)
    host_port, db = rest.rsplit('/', 1)
    host, port = (host_port.rsplit(':', 1) if ':' in host_port else (host_port, '5432'))
    kwargs = dict(host=host, port=int(port), database=db, user=user, password=password)
    if host not in ('localhost', '127.0.0.1'):
        kwargs['ssl_context'] = True
    return pg.Connection(**kwargs)


# Mapa de tipo pelo campo 'Tipo' da planilha
TIPO_MAP = {
    'Cliente':        'Cliente',
    'Interno':     'Interno',
    'Comunidade':  'Comunidade',
    'Treinamento': 'Treinamento',
}

def normalizar_tipo(t):
    if not t:
        return 'Outro'
    for k, v in TIPO_MAP.items():
        if k.lower() in str(t).lower():
            return v
    return 'Outro'


def main():
    conn = get_conn()

    wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb['Acompanhamento Ações SESMT']

    headers = None
    inseridos = 0
    erros = 0

    for row in ws.iter_rows(values_only=True):
        if headers is None:
            headers = row
            continue

        data_ev, evento, pessoas, obs, colab, cargo, contrato, tipo = row[:8]

        if not data_ev or not evento:
            continue

        if isinstance(data_ev, str):
            from datetime import date as d
            data_ev = d.fromisoformat(data_ev[:10])

        try:
            conn.run("""
                INSERT INTO acoes_sesmt
                    (data_evento, evento, pessoas_impactadas, observacoes,
                     colaborador_nome, cargo, contrato, tipo)
                VALUES (:dat, :ev, :pess, :obs, :col, :carg, :cont, :tipo)
                ON CONFLICT DO NOTHING
            """,
                dat=data_ev,
                ev=str(evento)[:1000] if evento else '',
                pess=int(pessoas) if pessoas else 0,
                obs=str(obs) if obs else None,
                col=str(colab) if colab else None,
                carg=str(cargo) if cargo else None,
                cont=str(contrato) if contrato else None,
                tipo=normalizar_tipo(tipo),
            )
            inseridos += 1
        except Exception as e:
            print(f"  Erro na linha: {e}")
            erros += 1

    conn.commit()
    conn.close()

    print(f"\n✅ Importação concluída: {inseridos} registros inseridos, {erros} erros.")


if __name__ == '__main__':
    main()
